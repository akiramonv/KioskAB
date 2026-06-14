from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


F_TITLE = font(34, True)
F_H = font(25, True)
F = font(22)
F_SMALL = font(18)


def wrap_text(text, max_chars):
    return "\n".join(textwrap.wrap(text, max_chars, break_long_words=False))


def centered_text(draw, box, text, fnt, fill=(24, 24, 24), spacing=6):
    lines = text.splitlines()
    heights = [draw.textbbox((0, 0), line, font=fnt)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total) / 2
    for line, h in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=fnt)
        x = box[0] + (box[2] - box[0] - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=fnt, fill=fill)
        y += h + spacing


def arrow(draw, start, end, fill=(38, 84, 124), width=4):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - direction * 18, y2 - 10), (x2 - direction * 18, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 10, y2 - direction * 18), (x2 + 10, y2 - direction * 18)]
    draw.polygon(pts, fill=fill)


def poly_arrow(draw, points, fill=(38, 84, 124), width=4):
    for a, b in zip(points, points[1:-1]):
        draw.line([a, b], fill=fill, width=width)
    arrow(draw, points[-2], points[-1], fill=fill, width=width)


def label(draw, xy, text, fnt=F_SMALL, fill=(38, 38, 38), max_chars=38):
    rendered = wrap_text(text, max_chars)
    bbox = draw.multiline_textbbox(xy, rendered, font=fnt, spacing=4)
    pad = 6
    draw.rectangle(
        (bbox[0] - pad, bbox[1] - pad, bbox[2] + pad, bbox[3] + pad),
        fill="white",
    )
    draw.multiline_text(xy, rendered, font=fnt, fill=fill, spacing=4)


def draw_context(path):
    img = Image.new("RGB", (2100, 1300), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 2099, 1299), outline=(180, 190, 200), width=3)
    d.text((70, 45), "IDEF0 TO-BE. Контекстная диаграмма A-0", font=F_TITLE, fill=(16, 54, 88))

    box = (620, 420, 1510, 805)
    d.rounded_rectangle(box, radius=12, fill=(236, 246, 241), outline=(30, 95, 80), width=5)
    centered_text(
        d,
        box,
        "A-0\nОбеспечить самообслуживание клиента\nчерез информационно-платежный терминал\nс QR-оплатой",
        F_H,
    )

    arrow(d, (130, 565), (620, 565))
    arrow(d, (130, 675), (620, 675))
    label(d, (70, 500), "Запрос клиента на услугу")
    label(d, (70, 610), "Данные плательщика, сумма, выбранная услуга")

    controls = [
        ("Регламенты банка и Abilling", 655),
        ("Правила ролей и безопасности", 900),
        ("Тарифы, комиссии, справочники", 1160),
        ("Требования к UI, надежности и журналированию", 1390),
    ]
    for text, x in controls:
        arrow(d, (x, 235), (x, 420))
        label(d, (x - 105, 155), text, max_chars=23)

    mechanisms = [
        ("Терминал самообслуживания, JavaFX UI", 660),
        ("KioskApp: контроллеры, ApiService, ApiClient", 970),
        ("ApiAB, PostgreSQL, Abilling / QR-сервис", 1320),
    ]
    for text, x in mechanisms:
        arrow(d, (x, 1055), (x, 805))
        label(d, (x - 125, 1090), text, max_chars=25)

    arrow(d, (1510, 520), (1980, 520))
    arrow(d, (1510, 625), (1980, 625))
    arrow(d, (1510, 730), (1980, 730))
    label(d, (1600, 455), "Сформированный платеж и QR-код", max_chars=32)
    label(d, (1600, 590), "Статус операции, чек/результат оплаты", max_chars=32)
    label(d, (1600, 695), "Журналы, отчеты, обновленные справочники", max_chars=32)

    d.text((70, 1215), "Узел: A-0    Точка зрения: клиент, администратор, банк    Цель: показать целевой процесс после внедрения терминала", font=F_SMALL, fill=(80, 80, 80))
    img.save(path)


def draw_decomposition(path):
    img = Image.new("RGB", (2400, 1500), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 2399, 1499), outline=(180, 190, 200), width=3)
    d.text((70, 45), "IDEF0 TO-BE. Декомпозиция A0", font=F_TITLE, fill=(16, 54, 88))

    boxes = {
        "A1": (210, 260, 610, 430, "A1\nАвторизовать пользователя\nи определить сессию"),
        "A2": (760, 250, 1160, 430, "A2\nЗагрузить и представить\nкаталог услуг"),
        "A3": (1320, 265, 1720, 455, "A3\nСформировать корзину\nи параметры платежа"),
        "A4": (480, 690, 890, 900, "A4\nСоздать платеж\nи отобразить QR-код"),
        "A5": (1050, 700, 1460, 910, "A5\nКонтролировать статус\nи зафиксировать результат"),
        "A6": (1640, 690, 2060, 910, "A6\nАдминистрировать\nсправочники и операции"),
    }
    for key, (x1, y1, x2, y2, text) in boxes.items():
        d.rounded_rectangle((x1, y1, x2, y2), radius=10, fill=(238, 248, 244), outline=(30, 95, 80), width=4)
        centered_text(d, (x1, y1, x2, y2), text, F_H if key in ("A1", "A2", "A3") else F, spacing=5)

    label(d, (70, 250), "Учетные данные\nпользователя", max_chars=24)
    arrow(d, (140, 320), (210, 320))
    label(d, (70, 520), "Запрос услуги,\nпоиск, категория", max_chars=24)
    poly_arrow(d, [(145, 570), (700, 570), (700, 345), (760, 345)])
    label(d, (70, 745), "Данные плательщика,\nсумма, количество", max_chars=25)
    arrow(d, (155, 790), (480, 790))

    arrow(d, (610, 340), (760, 340))
    label(d, (628, 285), "сессия, роль,\nпровайдер, специализации", max_chars=22)
    arrow(d, (1160, 345), (1320, 345))
    label(d, (1180, 285), "выбранная услуга\nи реквизиты", max_chars=22)
    poly_arrow(d, [(1520, 455), (1520, 555), (705, 555), (705, 690)])
    label(d, (940, 510), "подтвержденная корзина\nи итоговая сумма", max_chars=30)
    arrow(d, (890, 800), (1050, 800))
    label(d, (910, 740), "платеж, QR,\nкомиссия", max_chars=20)
    arrow(d, (1460, 805), (1640, 805))
    label(d, (1480, 745), "статусы, логи,\nплатежи", max_chars=20)

    for x in [330, 930, 1510]:
        arrow(d, (x, 150), (x, 250))
    label(d, (280, 105), "Политика доступа", max_chars=18)
    label(d, (840, 105), "Справочники, тарифы", max_chars=18)
    label(d, (1430, 105), "Правила оплаты", max_chars=18)
    label(d, (600, 1010), "JavaFX, SessionManager,\nApiService, ApiClient", max_chars=26)
    label(d, (1130, 1010), "ApiAB, PostgreSQL,\nAbilling QR", max_chars=26)
    label(d, (1745, 1010), "Администратор,\nжурналирование", max_chars=26)
    for x in [690, 1260, 1850]:
        arrow(d, (x, 1000), (x, 910))

    arrow(d, (2060, 750), (2300, 750))
    label(d, (2115, 690), "Обновленные справочники,\nнастройки, роли", max_chars=28)
    arrow(d, (1460, 875), (2300, 875))
    label(d, (1585, 930), "Итоговый статус, результат оплаты,\nзаписи журнала и отчеты", max_chars=36)

    d.text((70, 1415), "Узел: A0    Родитель: A-0    Модель показывает целевой порядок работы киоска после автоматизации", font=F_SMALL, fill=(80, 80, 80))
    img.save(path)


SECTIONS = [
    {
        "id": "A1",
        "name": "Авторизовать пользователя и определить контекст сессии",
        "purpose": "Проверить учетные данные, определить роль пользователя, организацию и доступные специализации.",
        "inputs": "email, пароль, данные пользователя из ApiAB",
        "controls": "правила авторизации, модель ролей user/admin/superAdmin, политика ограничения доступа",
        "outputs": "активная сессия, роль, провайдер, список специализаций, доступность административной панели",
        "mechanisms": "LoginController, ApiService.authenticate(), SessionManager, ApiAB /api/users и /api/roles",
        "sub": [
            "A1.1 Принять email и пароль на форме входа.",
            "A1.2 Проверить пользователя по справочнику ApiAB.",
            "A1.3 Получить роли, специализации и организацию пользователя.",
            "A1.4 Сохранить контекст в SessionManager.",
            "A1.5 Открыть главный экран и показать админ-панель только при наличии роли admin/superAdmin.",
        ],
    },
    {
        "id": "A2",
        "name": "Загрузить и представить каталог услуг",
        "purpose": "Предоставить клиенту актуальный список услуг, категорий и провайдеров с учетом его прав и специализации.",
        "inputs": "запрос клиента, активная сессия, справочники услуг, категорий и провайдеров",
        "controls": "категории, специализации пользователя, правила фильтрации по провайдеру, требования к сенсорному UI",
        "outputs": "экран каталога, карточки услуг, популярные услуги, выбранная услуга",
        "mechanisms": "MainController, ServiceCard, CategoryButton, ApiService.getAllServices(), getServicesByUser(), getAllCategories()",
        "sub": [
            "A2.1 Получить справочники провайдеров и категорий.",
            "A2.2 Загрузить все услуги либо услуги пользователя по специализациям.",
            "A2.3 Построить список категорий и карточки услуг.",
            "A2.4 Применить поиск по названию услуги, провайдеру и категории.",
            "A2.5 Открыть карточку детали услуги для дальнейшего формирования платежа.",
        ],
    },
    {
        "id": "A3",
        "name": "Сформировать корзину и параметры платежа",
        "purpose": "Преобразовать выбранную услугу и ввод клиента в корректные параметры платежной операции.",
        "inputs": "выбранная услуга, сумма, количество, данные плательщика, текущий провайдер",
        "controls": "проверка суммы больше 0, обязательность ФИО, связь услуги с провайдером и счетом, правила комиссии",
        "outputs": "корзина, итоговая сумма, подтвержденные реквизиты платежа",
        "mechanisms": "ServiceDetailController, CartItem, MainController.groupCartByProvider(), справочники провайдеров и комиссий",
        "sub": [
            "A3.1 Отобразить детали услуги: категорию, провайдера, счет и комиссию.",
            "A3.2 Проверить сумму и количество; для услуги с фиксированной ценой заблокировать изменение суммы.",
            "A3.3 Определить провайдера по provId либо названию из DTO.",
            "A3.4 Добавить позицию в корзину и пересчитать итог.",
            "A3.5 Получить ФИО и ИНН плательщика на этапе подтверждения.",
        ],
    },
    {
        "id": "A4",
        "name": "Создать платеж и отобразить QR-код",
        "purpose": "Передать подтвержденные параметры в ApiAB/Abilling, получить платеж и показать QR-код клиенту.",
        "inputs": "подтвержденная корзина, итоговая сумма, provId, ФИО, ИНН",
        "controls": "формат PaymentRequest, статус processing, правила расчета комиссии, доступность API",
        "outputs": "созданный платеж, QR-код, QR-ссылка, сумма, комиссия и сумма к оплате",
        "mechanisms": "ApiService.createPayment(), ApiClient.post('/api/payments'), PaymentController, Abilling / QR-сервис",
        "sub": [
            "A4.1 Сгруппировать позиции корзины по провайдеру для проверки корректности.",
            "A4.2 Рассчитать общую сумму платежа.",
            "A4.3 Отправить запрос создания платежа в ApiAB.",
            "A4.4 Получить id платежа, комиссию, итоговую сумму и QR-данные.",
            "A4.5 Отобразить QR-код и данные платежа на экране терминала.",
        ],
    },
    {
        "id": "A5",
        "name": "Контролировать статус и зафиксировать результат операции",
        "purpose": "Показать клиенту результат оплаты и оставить trace операций для администратора и сопровождения.",
        "inputs": "созданный платеж, QR-данные, ответ платежного сервиса, действия пользователя",
        "controls": "допустимые статусы processing/success/cancel, требования журналирования, правила обработки ошибок",
        "outputs": "финальный/текущий статус, сообщение клиенту, журнал событий, запись платежа для админки",
        "mechanisms": "Payment model, ApiService.getPaymentById(), updatePaymentStatus(), AppLogger, admin filters",
        "sub": [
            "A5.1 Получать или обновлять статус платежа через API.",
            "A5.2 Отображать понятный статус: ожидает оплаты, оплачен, отменен или ошибка.",
            "A5.3 Фиксировать создание платежа, ошибки и смену статуса в журнале.",
            "A5.4 Передавать платеж в административный контур для просмотра и фильтрации.",
            "A5.5 При сбое соединения сохранить управляемое состояние интерфейса без аварийного завершения.",
        ],
    },
    {
        "id": "A6",
        "name": "Администрировать справочники и операции",
        "purpose": "Обеспечить настройку данных системы и контроль операций без прямого вмешательства в клиентский сценарий.",
        "inputs": "действия администратора, справочники, платежи, пользователи, роли, специализации",
        "controls": "разграничение admin/superAdmin, привязка администратора к провайдеру, правила CRUD и целостности данных",
        "outputs": "обновленные пользователи, услуги, категории, провайдеры, специализации, статусы платежей, отчеты",
        "mechanisms": "AdminController, ApiService CRUD endpoints, ApiAB controllers, PostgreSQL, журналирование",
        "sub": [
            "A6.1 Открыть админ-панель только авторизованному администратору.",
            "A6.2 Загрузить раздел: пользователи, услуги, категории, провайдеры, специализации или платежи.",
            "A6.3 Ограничить данные провайдером для обычного admin; superAdmin видит все.",
            "A6.4 Создавать, редактировать и удалять справочники через ApiAB.",
            "A6.5 Фильтровать платежи по статусу, сумме и дате; изменять статус при необходимости.",
        ],
    },
]


def build_markdown(path):
    lines = []
    lines.append("# IDEF0 TO-BE модель системы информационно-платежного терминала")
    lines.append("")
    lines.append("Модель описывает целевое состояние системы после внедрения JavaFX-терминала KioskApp и API Emulator ApiAB. Точка зрения модели: клиент терминала, администратор организации и банк, контролирующий каталог услуг, платежи и учет операций.")
    lines.append("")
    lines.append("## Контекст A-0")
    lines.append("")
    lines.append("![IDEF0 TO-BE A-0](idef0_to_be_context.png)")
    lines.append("")
    lines.append("| Тип стрелки IDEF0 | Состав |")
    lines.append("|---|---|")
    lines.append("| Входы | запрос клиента на услугу; выбранная услуга; сумма; количество; ФИО и ИНН плательщика; учетные данные пользователя |")
    lines.append("| Управления | регламенты банка и Abilling; правила ролей и безопасности; тарифы и комиссии; справочники услуг, категорий, провайдеров и счетов; требования к надежности, UI и журналированию |")
    lines.append("| Выходы | созданный платеж; QR-код и QR-ссылка; статус операции; подтверждение для клиента; записи журнала; обновленные справочники и отчеты |")
    lines.append("| Механизмы | терминал самообслуживания; JavaFX-приложение KioskApp; ApiAB; PostgreSQL; Abilling/QR-сервис; администратор; сеть/VPN/TLS |")
    lines.append("")
    lines.append("## Декомпозиция A0")
    lines.append("")
    lines.append("![IDEF0 TO-BE A0](idef0_to_be_decomposition_a0.png)")
    lines.append("")
    lines.append("Главная функция A0 раскладывается на шесть связанных работ. Сначала система определяет пользователя и контекст доступа, затем загружает каталог услуг, формирует корзину, создает платеж, отображает QR-код, контролирует статус и передает операцию в административный контур.")
    lines.append("")
    lines.append("| Блок | Назначение | Основной результат |")
    lines.append("|---|---|---|")
    for s in SECTIONS:
        lines.append(f"| {s['id']} | {s['name']} | {s['outputs']} |")
    lines.append("")
    lines.append("## ICOM-описание работ A0")
    lines.append("")
    for s in SECTIONS:
        lines.append(f"### {s['id']}. {s['name']}")
        lines.append("")
        lines.append(s["purpose"])
        lines.append("")
        lines.append("| ICOM | Описание |")
        lines.append("|---|---|")
        lines.append(f"| Входы | {s['inputs']} |")
        lines.append(f"| Управления | {s['controls']} |")
        lines.append(f"| Выходы | {s['outputs']} |")
        lines.append(f"| Механизмы | {s['mechanisms']} |")
        lines.append("")
        lines.append("Декомпозиция:")
        lines.extend(f"- {item}" for item in s["sub"])
        lines.append("")
    lines.append("## Логика потоков между работами")
    lines.append("")
    flows = [
        "A1 -> A2: активная сессия, роль, провайдер и специализации определяют, какие услуги будут загружены.",
        "A2 -> A3: выбранная услуга передается в окно деталей, где клиент вводит сумму и количество.",
        "A3 -> A4: подтвержденная корзина, провайдер, ФИО и ИНН преобразуются в запрос создания платежа.",
        "A4 -> A5: созданный платеж с QR-кодом передается на экран оплаты и в контур контроля статуса.",
        "A5 -> A6: платеж, статус и журнал доступны администратору для контроля и анализа.",
        "A6 -> A2/A3/A4: обновленные справочники, провайдеры, счета, комиссии и специализации влияют на каталог и корректность платежа.",
    ]
    lines.extend(f"- {flow}" for flow in flows)
    lines.append("")
    lines.append("## Текст для вставки в пояснительную записку")
    lines.append("")
    lines.append("IDEF0-модель TO-BE отражает целевой процесс обслуживания клиента после внедрения стационарного информационно-платежного терминала. В отличие от состояния AS-IS, где значительная часть действий выполняется оператором, в модели TO-BE клиент самостоятельно выбирает услугу, вводит параметры платежа, подтверждает операцию и оплачивает ее по QR-коду. Администратор при этом не участвует в каждой клиентской операции, а выполняет настройку справочников, управление пользователями и контроль платежей через отдельную административную панель.")
    lines.append("")
    lines.append("Контекстная диаграмма A-0 показывает, что входами процесса являются запрос клиента, выбранная услуга, данные плательщика и учетные данные пользователя. Управляющими воздействиями выступают регламенты банка, правила безопасности, роли пользователей, тарифы, комиссии и требования к надежности системы. Механизмами выполнения процесса являются терминал самообслуживания, JavaFX-приложение KioskApp, API Emulator ApiAB, база данных PostgreSQL и платежная платформа Abilling. На выходе формируются платеж, QR-код, статус операции, подтверждение для клиента, журнал событий и данные для административного контроля.")
    lines.append("")
    lines.append("Декомпозиция A0 включает шесть основных функций: авторизацию пользователя, загрузку каталога услуг, формирование корзины, создание платежа и QR-кода, контроль статуса операции и администрирование справочников. Такая структура соответствует фактической архитектуре проекта: уровень представления реализован JavaFX-контроллерами, бизнес-логика сосредоточена в ApiService, HTTP-взаимодействие выполняет ApiClient, а состояние пользователя хранится в SessionManager. Благодаря этому целевой процесс является управляемым, масштабируемым и пригодным для эксплуатации на стационарном терминале.")
    path.write_text("\n".join(lines), encoding="utf-8")


def set_cell(cell, text, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.bold = bold


def build_docx(path, context_img, decomp_img):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("IDEF0 TO-BE модель системы информационно-платежного терминала")
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = RGBColor(16, 54, 88)

    p = doc.add_paragraph(
        "Модель описывает целевое состояние системы после внедрения JavaFX-терминала KioskApp и API Emulator ApiAB. "
        "Точка зрения модели: клиент терминала, администратор организации и банк."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_heading("Контекстная диаграмма A-0", level=1)
    doc.add_picture(str(context_img), width=Inches(7.1))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph("Рисунок 1 - IDEF0 TO-BE, контекстная диаграмма A-0")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_cell(table.rows[0].cells[0], "Тип стрелки", True)
    set_cell(table.rows[0].cells[1], "Состав", True)
    context_rows = [
        ("Входы", "запрос клиента; выбранная услуга; сумма; количество; ФИО и ИНН; учетные данные"),
        ("Управления", "регламенты банка и Abilling; роли; безопасность; тарифы; комиссии; справочники; требования к UI и журналированию"),
        ("Выходы", "созданный платеж; QR-код; статус операции; подтверждение; журнал; данные для админ-контроля"),
        ("Механизмы", "терминал; KioskApp; ApiAB; PostgreSQL; Abilling/QR-сервис; администратор; сеть/VPN/TLS"),
    ]
    for left, right in context_rows:
        row = table.add_row().cells
        set_cell(row[0], left, True)
        set_cell(row[1], right)

    doc.add_heading("Декомпозиция A0", level=1)
    doc.add_picture(str(decomp_img), width=Inches(7.1))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph("Рисунок 2 - IDEF0 TO-BE, декомпозиция A0")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Главная функция A0 раскладывается на шесть связанных работ: авторизация, загрузка каталога услуг, "
        "формирование корзины, создание платежа и QR-кода, контроль статуса и администрирование."
    )

    for s in SECTIONS:
        doc.add_heading(f"{s['id']}. {s['name']}", level=2)
        doc.add_paragraph(s["purpose"])
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        set_cell(t.rows[0].cells[0], "ICOM", True)
        set_cell(t.rows[0].cells[1], "Описание", True)
        for left, key in [("Входы", "inputs"), ("Управления", "controls"), ("Выходы", "outputs"), ("Механизмы", "mechanisms")]:
            row = t.add_row().cells
            set_cell(row[0], left, True)
            set_cell(row[1], s[key])
        doc.add_paragraph("Декомпозиция:")
        for item in s["sub"]:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Описание для пояснительной записки", level=1)
    for text in [
        "IDEF0-модель TO-BE отражает целевой процесс обслуживания клиента после внедрения стационарного информационно-платежного терминала. Клиент самостоятельно выбирает услугу, вводит параметры платежа, подтверждает операцию и оплачивает ее по QR-коду. Администратор выполняет настройку справочников, управление пользователями и контроль платежей через отдельную административную панель.",
        "Контекстная диаграмма A-0 показывает, что входами процесса являются запрос клиента, выбранная услуга, данные плательщика и учетные данные пользователя. Управляющими воздействиями выступают регламенты банка, правила безопасности, роли пользователей, тарифы, комиссии и требования к надежности системы. Механизмами выполнения процесса являются терминал самообслуживания, JavaFX-приложение KioskApp, API Emulator ApiAB, база данных PostgreSQL и платежная платформа Abilling.",
        "Декомпозиция A0 соответствует фактической архитектуре проекта: уровень представления реализован JavaFX-контроллерами, бизнес-логика сосредоточена в ApiService, HTTP-взаимодействие выполняет ApiClient, а состояние пользователя хранится в SessionManager. Благодаря этому целевой процесс является управляемым, масштабируемым и пригодным для эксплуатации на стационарном терминале.",
    ]:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(path)


def main():
    OUT.mkdir(exist_ok=True)
    context = OUT / "idef0_to_be_context.png"
    decomp = OUT / "idef0_to_be_decomposition_a0.png"
    draw_context(context)
    draw_decomposition(decomp)
    build_markdown(OUT / "IDEF0_TO_BE_KioskApp.md")
    build_docx(OUT / "IDEF0_TO_BE_KioskApp.docx", context, decomp)


if __name__ == "__main__":
    main()
