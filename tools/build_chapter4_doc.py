from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Глава_4_Экспериментальный_раздел_KioskApp.docx"


def load_json(path: str):
    raw = ROOT / path
    try:
        return json.loads(raw.read_text(encoding="utf-8"))
    except UnicodeDecodeError:
        return json.loads(raw.read_text(encoding="utf-16"))


metrics = load_json("target/test-results/metrics.json")
manual = load_json("target/test-results/manual-tests.json")
fxml = load_json("target/test-results/fxml-validation.json")


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(12)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(width)
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(12)
        if row_i == 0:
            set_repeat_table_header(row)
            for cell in row.cells:
                set_cell_shading(cell, "E8EEF5")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def add_table(doc: Document, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    style_table(table, widths)
    doc.add_paragraph()
    return table


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_para(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_bullets(doc: Document, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def configure_doc(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 14)]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.5

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ГЛАВА 4: ЭКСПЕРИМЕНТАЛЬНЫЙ РАЗДЕЛ")
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(16)

    add_heading(doc, "4.1. План тестирования", 1)
    add_para(
        doc,
        "Данный раздел описывает экспериментальную проверку JavaFX-приложения KioskApp, "
        "предназначенного для работы киоска самообслуживания с API Emulator. Цель тестирования - "
        "подтвердить корректность основных пользовательских сценариев, устойчивость модулей, "
        "правильность обмена данными с API и качество программного кода. Нагрузочное тестирование "
        "в соответствии с заданием не проводилось."
    )
    add_para(
        doc,
        "В план тестирования включены: тестирование методом черного ящика, тестирование методом "
        "белого ящика, модульное тестирование, интеграционное тестирование, системная проверка "
        "сборки и ресурсов, статический анализ кода, профилирование тестового прогона и расчет "
        "метрик кода."
    )

    add_heading(doc, "4.2. Объекты тестирования", 1)
    add_bullets(
        doc,
        [
            "Главный экран киоска: загрузка услуг, категории, поиск, фильтрация, популярные услуги.",
            "Авторизация: проверка email и пароля, загрузка ролей, специализаций и провайдера пользователя.",
            "Корзина и оплата: добавление услуги, расчет суммы, группировка по провайдеру, создание платежа.",
            "Административная панель: управление пользователями, услугами, категориями, провайдерами, специализациями и платежами.",
            "Интеграционный слой: ApiClient и ApiService, REST-запросы к API Emulator, обработка success/data/message.",
            "Модели данных: User, Role, ProviderService, CartItem, Payment, Provider и справочные классы.",
            "FXML-ресурсы и CSS-тема JavaFX-приложения.",
        ],
    )

    add_heading(doc, "4.3. Тестовая среда и начальные условия", 1)
    add_table(
        doc,
        ["Параметр", "Фактическое значение", "Результат"],
        [
            ["Операционная среда", "Windows, PowerShell, рабочая папка kiosk-app", "Среда доступна"],
            ["Основной java в PATH", "Java 1.8.0_471", "Не подходит для проекта Java 17"],
            ["Дополнительный JDK", "OpenJDK 25, компиляция с --release 17", "Использован для проверки"],
            ["Maven CLI", "Команда mvn отсутствует", "Maven-тесты не запускались"],
            ["API Emulator", "localhost:7924 не отвечает", "Реальная API-интеграция ограничена"],
            ["Зависимости", "JavaFX 21.0.2, Jackson 2.16.1", "Доступны локально"],
        ],
        widths=[4.0, 7.0, 5.0],
    )

    add_heading(doc, "4.4. Тестирование методом черного ящика", 1)
    add_para(
        doc,
        "Тестирование методом черного ящика проводилось по внешнему поведению системы: входные данные, "
        "ожидаемые реакции интерфейса и API-слоя. Для сценариев, требующих настоящего API Emulator, "
        "результат отмечен как ограниченный средой, так как порт 7924 на момент проверки был недоступен."
    )
    add_table(
        doc,
        ["№", "Входные данные / действие", "Ожидаемый результат", "Фактический результат"],
        [
            [1, "Открытие главного экрана при доступном API", "Загружаются категории и услуги", "Проверено через сборку и FXML; реальный API недоступен"],
            [2, "Поиск услуги по части названия", "Список услуг фильтруется по названию, провайдеру или категории", "Логика обнаружена в MainController.applyFilters"],
            [3, "Выбор категории", "Отображаются услуги выбранной категории", "Проверено анализом ветви фильтрации"],
            [4, "Пустые email или пароль", "Показывается ошибка ввода", "Проверено анализом LoginController.onLogin"],
            [5, "Неверные учетные данные", "Показывается ошибка авторизации", "Проверено через сценарий authenticate"],
            [6, "Роль user", "Кнопка админ-панели скрыта", "Проверено модульным тестом ролей"],
            [7, "Роль admin", "Кнопка админ-панели доступна", "Проверено модульным тестом ролей"],
            [8, "Услуга с фиксированной ценой", "Сумма подставляется автоматически", "Проверено модульным тестом ProviderService"],
            [9, "Количество 0 в корзине", "Количество нормализуется до 1 или отклоняется на UI", "Проверено модульным тестом CartItem"],
            [10, "Создание платежа без провайдера", "Показывается ошибка невозможности создать платеж", "Проверено анализом MainController.onCheckout"],
            [11, "API возвращает HTTP 500", "Клиент выбрасывает исключение с message", "Проверено mock API"],
            [12, "FXML-файлы приложения", "Все контроллеры находятся и ресурсы читаются", "4 из 4 PASS"],
        ],
        widths=[1.0, 5.2, 5.2, 5.0],
    )

    add_heading(doc, "4.5. Тестирование методом белого ящика", 1)
    add_para(
        doc,
        "При тестировании методом белого ящика анализировались ветвления, условия и независимые пути "
        "в ключевых методах. Особое внимание уделено методам, от которых зависит корректность платежей, "
        "прав доступа и API-обмена."
    )
    add_table(
        doc,
        ["Метод", "Проверяемые пути", "Результат"],
        [
            ["Role.isAdmin / isSuperAdmin", "user, admin, superAdmin", "Пути проверены, результат корректный"],
            ["SessionManager.logout", "Очистка пользователя, провайдера, ролей и специализаций", "Обнаружен дефект при неизменяемом списке ролей"],
            ["CartItem.getTotal", "Округление суммы, умножение на quantity, quantity <= 0", "Пути проверены, результат корректный"],
            ["Payment.getTotal", "sum/fee заданы, sum/fee равны null", "Пути проверены, результат корректный"],
            ["Provider.matchesName", "shortName, fullName, toString, null", "Пути проверены, результат корректный"],
            ["ApiClient.send", "2xx success, HTTP 500, success=false", "Пути проверены через mock API"],
            ["MainController.applyFilters", "Без фильтров, категория, строка поиска", "Проверено анализом кода"],
            ["AdminController.matchesPaymentFilters", "Статус, минимальная/максимальная сумма, даты", "Проверено анализом ветвлений"],
        ],
        widths=[4.2, 7.0, 5.0],
    )

    add_heading(doc, "4.6. Модульное тестирование", 1)
    unit_rows = []
    for idx, result in enumerate(manual["results"], 1):
        status, text = result.split(": ", 1)
        unit_rows.append([idx, text, status])
    add_table(doc, ["№", "Тест", "Статус"], unit_rows, widths=[1.0, 11.5, 3.0])
    add_para(
        doc,
        f"Итог модульного прогона: успешно - {manual['passed']}, неуспешно - {manual['failed']}. "
        "Неуспешный тест связан не с тестовым стендом, а с найденной ошибкой в управлении состоянием сессии."
    )

    add_heading(doc, "4.7. Интеграционное тестирование", 1)
    add_para(
        doc,
        "Так как реальный API Emulator на localhost:7924 не отвечал, для проверки интеграционного слоя "
        "был поднят локальный mock HTTP-сервер на том же порту. Он имитировал успешный ответ /api/services, "
        "универсальный /api/emulator/execute и ошибочный HTTP 500."
    )
    add_table(
        doc,
        ["Интеграционный сценарий", "Ожидаемый результат", "Фактический результат"],
        [
            ["ApiService.getAllServices -> ApiClient.get -> mock /api/services", "Список услуг преобразуется в ProviderService", "PASS"],
            ["ApiService.executeOperation -> mock /api/emulator/execute", "Ответ success=true корректно читается", "PASS"],
            ["ApiClient.get -> mock /api/fail HTTP 500", "Исключение IllegalStateException с текстом ошибки", "PASS"],
            ["Подключение к настоящему API Emulator", "Порт 7924 отвечает", "BLOCKED: TcpTestSucceeded=False"],
        ],
        widths=[6.5, 6.0, 3.5],
    )

    add_heading(doc, "4.8. Системная проверка сборки и ресурсов", 1)
    add_table(
        doc,
        ["Проверка", "Команда / способ", "Результат"],
        [
            ["Maven", "mvn -version", "FAIL: Maven CLI не установлен"],
            ["Немодульная компиляция", "javac --release 17 без module-info.java", "PASS"],
            ["Модульная компиляция", "javac --release 17 с module-info.java", "PASS"],
            ["Количество class-файлов после модульной сборки", "target/manual-compile-module", "25 class-файлов"],
            ["FXML-валидация", "XML parse + проверка fx:controller", "4 из 4 PASS"],
            ["CSS-ресурс", "src/main/resources/css/styles.css", "Файл найден"],
        ],
        widths=[4.2, 6.5, 5.3],
    )

    add_heading(doc, "4.9. Статический анализ кода", 1)
    add_para(
        doc,
        "Статический анализ проводился поиском рискованных конструкций, проверкой структуры JavaFX-ресурсов "
        "и оценкой основных точек сложности. Критических синтаксических ошибок не выявлено, сборка через javac "
        "прошла успешно."
    )
    add_table(
        doc,
        ["Наблюдение", "Файл / область", "Оценка"],
        [
            ["BASE_URL жестко задан как http://localhost:7924", "ApiClient.java", "Средний риск: лучше вынести в конфигурацию"],
            ["Пароль хранится в модели и редактируется как обычная строка", "User.java, AdminController.java", "Высокий риск безопасности"],
            ["Фоновые операции запускаются через new Thread()", "LoginController, MainController, AdminController", "Средний риск: желательно ExecutorService/Task"],
            ["UI обновляется через Platform.runLater()", "Контроллеры JavaFX", "Корректная практика для JavaFX"],
            ["TODO/FIXME не обнаружены", "Все Java-файлы", "PASS"],
            ["FXML-контроллеры существуют", "resources/fxml", "PASS"],
            ["Самые крупные классы: AdminController и MainController", "ui/controllers", "Риск сопровождения из-за размера"],
        ],
        widths=[6.0, 5.0, 5.0],
    )

    add_heading(doc, "4.10. Профилирование", 1)
    mem_delta = manual["usedMemoryAfterBytes"] - manual["usedMemoryBeforeBytes"]
    add_para(
        doc,
        "Профилирование выполнялось на тестовом стенде без нагрузочного режима. Измерялись время выполнения "
        "модульно-интеграционного прогона и изменение используемой памяти JVM."
    )
    add_table(
        doc,
        ["Показатель", "Значение"],
        [
            ["Время выполнения тестового стенда", f"{manual['durationMs']} мс"],
            ["Используемая память до запуска", f"{manual['usedMemoryBeforeBytes']:,} байт".replace(",", " ")],
            ["Используемая память после запуска", f"{manual['usedMemoryAfterBytes']:,} байт".replace(",", " ")],
            ["Изменение используемой памяти", f"{mem_delta:,} байт".replace(",", " ")],
            ["Количество проверок в стенде", str(manual["passed"] + manual["failed"])],
        ],
        widths=[7.0, 5.0],
    )

    add_heading(doc, "4.11. Метрики кода", 1)
    add_table(
        doc,
        ["Метрика", "Значение"],
        [
            ["Java-файлов", metrics["java_files"]],
            ["FXML-файлов", metrics["fxml_files"]],
            ["CSS-файлов", metrics["css_files"]],
            ["Физических строк Java", metrics["physical_lines_java"]],
            ["Строк кода Java", metrics["code_lines_java"]],
            ["Пустых строк", metrics["blank_lines_java"]],
            ["Комментариев", metrics["comment_lines_java"]],
            ["Классов/интерфейсов/enum", metrics["classes_interfaces_enums"]],
            ["Методов приблизительно", metrics["methods_approx"]],
            ["Ветвлений приблизительно", metrics["branches_approx"]],
            ["Суммарная приблизительная цикломатическая сложность", metrics["approx_cyclomatic_total"]],
            ["TODO/FIXME", metrics["todos"]],
        ],
        widths=[8.0, 4.0],
    )
    top_rows = [
        [item["file"].replace("src\\main\\java\\", ""), item["lines"], item["methods"], item["branches"]]
        for item in metrics["files_top_by_lines"][:6]
    ]
    add_table(doc, ["Файл", "Строк", "Методов", "Ветвлений"], top_rows, widths=[8.5, 2.0, 2.0, 2.5])

    add_heading(doc, "4.12. Обнаруженные дефекты и рекомендации", 1)
    add_table(
        doc,
        ["№", "Дефект / риск", "Приоритет", "Рекомендация"],
        [
            [1, "SessionManager.logout() падает с UnsupportedOperationException, если роли переданы неизменяемым списком", "Высокий", "В setCurrentUserRoles и setCurrentUserSpecializations сохранять new ArrayList<>(...)"],
            [2, "В PATH доступна Java 8, а проект требует Java 17+", "Высокий", "Настроить JAVA_HOME/JDK 17+ в IntelliJ и терминале"],
            [3, "Maven CLI отсутствует", "Средний", "Установить Maven или использовать Maven Wrapper"],
            [4, "API Emulator недоступен на localhost:7924", "Высокий", "Запускать emulator перед интеграционными и приемочными тестами"],
            [5, "Пароль хранится в User как открытая строка", "Высокий", "Использовать хеширование и не показывать пароль в админ-форме"],
            [6, "BASE_URL жестко зашит в ApiClient", "Средний", "Вынести адрес API в properties/env-параметр"],
            [7, "AdminController содержит 874 строки и 187 ветвлений", "Средний", "Разделить CRUD, фильтры и диалоги на отдельные классы"],
        ],
        widths=[1.0, 7.0, 2.2, 5.8],
    )

    add_heading(doc, "4.13. Вывод по экспериментальному разделу", 1)
    add_para(
        doc,
        "В ходе экспериментальной проверки проект KioskApp был протестирован несколькими методами. "
        "Сборка исходного кода с целевым уровнем Java 17 прошла успешно при использовании JDK 25 и локальных "
        "зависимостей. FXML-ресурсы корректны, основные модели и интеграционный слой прошли модульные и "
        "mock-интеграционные проверки. При этом выявлены ограничения среды: Maven CLI отсутствует, основной "
        "java в PATH имеет версию 1.8.0_471, а реальный API Emulator на localhost:7924 недоступен."
    )
    add_para(
        doc,
        "Главный обнаруженный дефект связан с очисткой сессии при неизменяемом списке ролей. После его "
        "устранения рекомендуется повторить тестовый стенд, затем провести приемочное тестирование с реально "
        "запущенным API Emulator. Нагрузочное тестирование в рамках данной главы не выполнялось."
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
